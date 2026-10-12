"""農事作業紀錄表匯出（PDF 與 Excel）。

和診斷書同一個做法：API 先把資料整理成快照字典，這裡只負責組表與輸出，不持有 DB session。
對齊縣市補助表單時只換 templates/farm_log_export.html，資料準備不動。

幾條刻意的規則（改之前先看理由）：
* 欄位完整度用「目前」的規則版本重算，只拿來在匯出時提醒；紀錄存檔當時的判定不回寫。
* 安全採收期核對只標示、不擋匯出：紀錄表要照實反映發生過的事，能不能採收由農民與植物醫生判斷。
* 施藥的安全採收期只取處方或藥劑登記快照；手打的藥名沒有依據，一律「無法判斷」，不猜、不填 0。
"""
from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from jinja2 import Environment, FileSystemLoader, StrictUndefined, select_autoescape
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

from app.services.farm_log import COMPLIANCE_RULE_VERSION, KIND_LABELS, as_utc, compliance, default_title, plain_number

TAIPEI = ZoneInfo("Asia/Taipei")
TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"
# 同田區往前找多久的施藥。核准的安全採收期很少超過數十天，一年足以涵蓋，又不會把好幾期前的作物拉進來。
PHI_LOOKBACK = timedelta(days=365)
# 附錄照片縮成長邊 1000px 的 JPEG：印在 A4 半頁上夠清楚，上百張照片的 PDF 也不至於大到打不開。
APPENDIX_IMAGE_PX = 1000
MATERIAL_KINDS = ("spray", "fertilize")
FIELD_LABELS = {
    "material_name": "資材品名",
    "material_supplier": "供應者",
    "material_lot": "批號",
    "amount_value": "用量",
    "amount_unit": "用量單位",
    "area_text": "面積",
    "crop_text": "品項",
    "yield_value": "產量",
    "yield_unit": "產量單位",
}
PHI_RESULTS = {"ok": "符合", "violation": "未達安全採收期", "unknown": "無法判斷", "no_spray": "無施藥紀錄"}
PHI_SOURCES = {"prescription": "處方", "registration": "藥劑登記"}


@dataclass(frozen=True)
class ExportOptions:
    photos: bool = True
    gps: bool = True
    phi: bool = True


# ---- 日期與文字 ----

def roc_day(value: date) -> str:
    return f"{value.year - 1911} 年 {value.month} 月 {value.day} 日"


def period_label(start: date, end: date) -> str:
    if start.year == end.year:
        return f"{roc_day(start)} 至 {end.month} 月 {end.day} 日"
    return f"{roc_day(start)} 至 {roc_day(end)}"


def short_day(value: date, with_year: bool = False) -> str:
    text = f"{value.month}/{value.day:02d}"
    return f"{value.year - 1911}/{text}" if with_year else text


def roc_datetime(value: datetime) -> str:
    local = as_utc(value).astimezone(TAIPEI)
    return f"{local.year - 1911}/{local.month:02d}/{local.day:02d} {local:%H:%M}"


def field_labels(names) -> list[str]:
    names = list(names)
    # 用量與用量單位同時缺只說一次「用量」，農民看到的是一個欄位不是兩個。
    return [FIELD_LABELS.get(name, name) for name in names
            if not (name.endswith("_unit") and name.replace("_unit", "_value") in names)]


def number_text(value: float) -> str:
    return f"{value:,.3f}".rstrip("0").rstrip(".")


def join_text(*values: str) -> str:
    return " ".join(value.strip() for value in values if value and value.strip())


# ---- 安全採收期核對 ----

def _spray_ref(spray: dict) -> dict:
    return {key: spray.get(key) for key in ("id", "done_on", "material_name", "phi_days", "phi_source")}


def _same_field(a: dict, b: dict) -> bool:
    # 兩邊都連到田區主檔才用編號比；否則退回名稱快照（寫法不同會對不到，這是名稱比對的已知限制）。
    if a.get("field_id") and b.get("field_id"):
        return a["field_id"] == b["field_id"]
    return a["field_text"] == b["field_text"]


def phi_checks(harvests: list[dict], sprays: list[dict]) -> list[dict]:
    """每筆採收對同田區、同一期規劃的施藥核對安全採收期。

    看的是「限制最嚴」的那次施藥（施藥日＋安全採收期最晚），不是最近一次：
    較早噴的藥安全採收期若比較長，可能比昨天噴的更晚才能採。
    有任何一次施藥的安全採收期未載明，就不能說「符合」，只能說無法判斷。
    """
    results = []
    for harvest in harvests:
        day = harvest["done_on"]
        related = [
            spray for spray in sprays
            if _same_field(spray, harvest) and day - PHI_LOOKBACK <= spray["done_on"] <= day
            # 同田區但屬於不同期規劃，是上一期作物的施藥，跟這次採收的作物無關。
            and not (spray.get("plan_id") and harvest.get("plan_id") and spray["plan_id"] != harvest["plan_id"])
        ]
        known = [spray for spray in related if spray["phi_days"] is not None]
        unknown = [spray for spray in related if spray["phi_days"] is None]
        binding = max(known, key=lambda s: (s["done_on"] + timedelta(days=s["phi_days"]), s["done_on"]), default=None)
        latest = max(related, key=lambda s: s["done_on"], default=None)
        entry = dict(harvest_id=harvest["id"], harvest_on=day, harvest_label=short_day(day), field_text=harvest["field_text"],
                     crop_text=harvest.get("crop_text") or "", days_since_last=(day - latest["done_on"]).days if latest else None,
                     spray=None, safe_from=None)
        if not related:
            entry.update(status="no_spray", message="回溯一年內，同田區沒有施藥紀錄")
        elif binding and (day - binding["done_on"]).days < binding["phi_days"]:
            safe = binding["done_on"] + timedelta(days=binding["phi_days"])
            entry.update(status="violation", spray=_spray_ref(binding), safe_from=safe, message=(
                f"距 {short_day(binding['done_on'])} 施用「{binding['material_name']}」僅 {(day - binding['done_on']).days} 天，"
                f"未達安全採收期 {binding['phi_days']} 天（{short_day(safe)} 起可採）"))
        elif unknown:
            pick = max(unknown, key=lambda s: s["done_on"])
            entry.update(status="unknown", spray=_spray_ref(pick), message=(
                f"{short_day(pick['done_on'])} 施用「{pick['material_name']}」沒有帶入處方或藥劑登記，安全採收期未載明，無法判斷"))
        else:
            entry.update(status="ok", spray=_spray_ref(binding), safe_from=binding["done_on"] + timedelta(days=binding["phi_days"]),
                         message=(f"距 {short_day(binding['done_on'])} 施藥 {(day - binding['done_on']).days} 天，"
                                  f"已過安全採收期 {binding['phi_days']} 天"))
        entry["result_label"] = PHI_RESULTS[entry["status"]]
        results.append(entry)
    return results


# ---- 組表 ----

def _title_detail(log: dict) -> str:
    """標題去掉「類型 · 」前綴；和依資材自動產生的標題相同時回空字串（資材欄已經寫了）。"""
    label = KIND_LABELS[log["kind"]]
    title = log["title"].strip()
    if not title or title == label or title == default_title(log["kind"], material_name=log["material_name"], dilution=log["dilution"]):
        return ""
    return title[len(label) + 3:] if title.startswith(f"{label} · ") else title


def _material_cell(log: dict, missing: list[str]) -> tuple[str, str]:
    values = [log[key].strip() for key in ("material_name", "material_supplier", "material_lot")]
    pending = field_labels(key for key in ("material_name", "material_supplier", "material_lot") if key in missing)
    if not any(values) and not pending:
        return "—", ""
    text = "／".join(value for value in values if value)
    return text, (f"{'／' if text else ''}（{'、'.join(pending)}待補）" if pending else "")


def _amount_cell(log: dict, missing: list[str]) -> tuple[str, str]:
    keys = ("yield_value", "yield_unit") if log["kind"] == "harvest" else ("amount_value", "amount_unit")
    text = join_text(*(log[key] for key in keys))
    lacking = [key for key in keys if key in missing]
    return text, (f"（{'、'.join(field_labels(lacking))}待補）" if lacking else "")


def _description(log: dict, phi_entry: dict | None, amount_used_labor: bool) -> str:
    kind = log["kind"]
    parts = []
    detail = _title_detail(log)
    if detail:
        parts.append(detail)
    pest = log["pest_text"].strip()
    if pest:
        parts.append(f"{pest}防治" if kind == "spray" else f"觀察：{pest}" if kind == "scout" else pest)
    ratio = log["dilution"].strip()
    if kind in MATERIAL_KINDS and ratio:
        parts.append(f"稀釋 {ratio}" if ratio.endswith("倍") else f"稀釋 {ratio} 倍")
    if log.get("case_no"):
        parts.append(f"依診斷書 {log['case_no']}")
    if kind == "spray":
        if log["phi_days"] is None:
            parts.append("安全採收期未載明")
        else:
            safe = log["done_on"] + timedelta(days=log["phi_days"])
            parts.append(f"安全採收期 {log['phi_days']} 天（{short_day(safe)} 起可採）")
    if kind == "harvest":
        if log["yield_destination"].strip():
            parts.append(f"出貨 {log['yield_destination'].strip()}")
        if phi_entry and phi_entry["days_since_last"] is not None:
            parts.append(f"距上次施藥 {phi_entry['days_since_last']} 天")
    if log["labor_text"].strip() and not amount_used_labor:
        parts.append(f"人力 {log['labor_text'].strip()}")
    if log["note"].strip():
        parts.append(log["note"].strip())
    return "；".join(parts) or "—"


def _photo_time(photo: dict) -> str:
    stamp = roc_datetime(photo["taken_at"])
    # 退回上傳時間的照片不能印得像拍攝時間：查核只認相機記錄的時間。
    return f"拍攝時間 {stamp}" if photo["taken_at_source"] == "exif" else f"上傳時間（無拍攝資訊）{stamp}"


def _photo_gps(photo: dict) -> str:
    if photo.get("lat") is None or photo.get("lng") is None:
        return "無 GPS 資訊"
    return f"GPS {photo['lat']:.5f}, {photo['lng']:.5f}"


def build_export(*, farm_name: str, operator: str, start: date, end: date, field: str | None, logs: list[dict],
                 sprays: list[dict], plans: dict[str, dict], options: ExportOptions, generated_on: date) -> dict:
    ordered = sorted(logs, key=lambda log: (log["done_on"], log["done_time"] is None, log["done_time"] or "",
                                            log["created_at"], log["id"]))
    with_year = start.year != end.year
    # 田區依第一次出現的先後排列，跟紀錄表的時間順序一致；依字碼排序對中文田區名稱沒有意義。
    fields = list(dict.fromkeys(log["field_text"] for log in ordered))
    phi = phi_checks([log for log in ordered if log["kind"] == "harvest"], sprays) if options.phi else []
    phi_by_harvest = {entry["harvest_id"]: entry for entry in phi}

    rows, appendix, incomplete, harvests = [], [], [], []
    totals: dict[str, float] = {}
    unparsed = 0
    rule_changed = 0
    for log in ordered:
        missing = compliance(log["kind"], log)["missing"]
        if (log.get("compliance") or {}).get("rule_version") != COMPLIANCE_RULE_VERSION:
            rule_changed += 1
        if missing:
            incomplete.append(dict(id=log["id"], done_on=log["done_on"], title=log["title"], missing=field_labels(missing)))

        numbers = []
        for photo in log["photos"]:
            numbers.append(len(appendix) + 1)
            appendix.append(dict(number=numbers[-1], path=photo.get("path"), image=None,
                                 log_date=short_day(log["done_on"], with_year), kind_label=KIND_LABELS[log["kind"]],
                                 field_text=log["field_text"], title=log["title"], caption=photo.get("caption") or "",
                                 time_text=_photo_time(photo), gps_text=_photo_gps(photo),
                                 taken_at=photo["taken_at"], taken_at_source=photo["taken_at_source"],
                                 lat=photo.get("lat"), lng=photo.get("lng")))
        if not numbers:
            photos_label = "—"
        elif not options.photos:
            photos_label = str(len(numbers))
        else:
            photos_label = str(numbers[0]) if len(numbers) == 1 else f"{numbers[0]}–{numbers[-1]}"

        material, material_pending = _material_cell(log, missing)
        amount, amount_pending = _amount_cell(log, missing)
        used_labor = not amount and not amount_pending and bool(log["labor_text"].strip())
        if used_labor:
            amount = log["labor_text"].strip()
        rows.append(dict(
            id=log["id"], log=log, date_label=short_day(log["done_on"], with_year), done_time=log["done_time"] or "",
            kind_label=KIND_LABELS[log["kind"]], field_text=log["field_text"], area_text=log["area_text"].strip(),
            material=material, material_pending=material_pending, amount=amount or ("" if amount_pending else "—"),
            amount_pending=amount_pending, description=_description(log, phi_by_harvest.get(log["id"]), used_labor),
            photos_label=photos_label, photo_count=len(numbers), missing=field_labels(missing), miss=bool(missing),
        ))

        if log["kind"] == "harvest":
            value = plain_number(log["yield_value"])
            unit = log["yield_unit"].strip()
            if log["yield_value"].strip():
                if value is None or not unit:
                    unparsed += 1
                else:
                    totals[unit] = totals.get(unit, 0.0) + value
            harvests.append(dict(date_label=short_day(log["done_on"], with_year), field_text=log["field_text"],
                                 item=" · ".join(v for v in (log["crop_text"].strip(), _title_detail(log)) if v) or "—",
                                 yield_text=join_text(log["yield_value"], log["yield_unit"]) or "—",
                                 yield_value=value if value is not None else log["yield_value"].strip(), yield_unit=unit,
                                 destination=log["yield_destination"].strip() or "—", done_on=log["done_on"]))

    areas = sorted({log["area_text"].strip() for log in ordered if log["area_text"].strip()}
                   | {plan["area"] for plan in plans.values() if plan.get("area") and (not field or plan["field_name"] == field)})
    crops = sorted({(log["crop_text"].strip(), (plans.get(log["plan_id"]) or {}).get("variety") or "")
                    for log in ordered if log["crop_text"].strip()})
    if field:
        field_label = field
    elif len(fields) > 3:
        field_label = f"{'、'.join(fields[:3])} 等 {len(fields)} 區"
    else:
        field_label = "、".join(fields) or "—"

    return dict(
        farm_name=farm_name or "—", operator=operator or "—", field_label=field_label,
        area_label="、".join(areas) or "未填", crop_label="、".join(f"{c} · {v}" if v else c for c, v in crops) or "未填",
        period_label=period_label(start, end), generated_label=roc_day(generated_on), start=start, end=end,
        rule_version=COMPLIANCE_RULE_VERSION, options=options, fields=fields, show_field_column=len(fields) > 1,
        rows=rows, harvests=harvests, harvest_totals=[dict(unit=unit, total=total, text=f"{number_text(total)} {unit}")
                                                      for unit, total in sorted(totals.items())],
        unparsed_yields=unparsed, phi=phi, appendix=appendix, any_missing=bool(incomplete),
        warnings=dict(incomplete=incomplete, rule_changed=rule_changed,
                      phi_violations=[entry for entry in phi if entry["status"] == "violation"],
                      phi_unknown=[entry for entry in phi if entry["status"] == "unknown"]),
        counts=dict(logs=len(rows), photos=len(appendix), harvests=len(harvests)),
    )


# ---- PDF ----

def appendix_image(path: Path | None, size: int = APPENDIX_IMAGE_PX) -> str | None:
    """附錄用的縮圖。依 EXIF 方向轉正（原檔保留方向標記，不轉就會印成橫躺）；檔案不在就回 None，版面標示缺檔。"""
    if path is None:
        return None
    try:
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source)
            image.thumbnail((size, size))
            if image.mode != "RGB":
                image = image.convert("RGB")
            buffer = io.BytesIO()
            image.save(buffer, "JPEG", quality=78, optimize=True)
    except (OSError, ValueError):
        return None
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def render_export_html(data: dict, font_url: str) -> str:
    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR), autoescape=select_autoescape(["html"]), undefined=StrictUndefined)
    return env.get_template("farm_log_export.html").render(d=data, font_url=font_url)


def render_export_pdf(data: dict) -> bytes:
    # 延後載入 Playwright：匯出預覽與 Excel 用不到瀏覽器。
    from app.services.consult_report import font_data, html_to_pdf

    if data["options"].photos:
        for item in data["appendix"]:
            item["image"] = appendix_image(item["path"])
    return html_to_pdf(render_export_html(data, font_data()), label="農事紀錄表", page_numbers=True)


# ---- Excel ----

HEAD_FILL = PatternFill("solid", fgColor="F2F0E9")
MISS_FILL = PatternFill("solid", fgColor="FDF7EA")
CRIT_FILL = PatternFill("solid", fgColor="FBEBE9")


def _put(sheet, row: int, column: int, value, fill=None):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        # 農民輸入的文字可能以 = 開頭；一律存成文字，Excel 開檔時不會被當成公式執行。
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if fill is not None:
        cell.fill = fill
    return cell


def _table(sheet, top: int, headers: list[str], rows: list[tuple[list, PatternFill | None]], widths: list[int]):
    for column, header in enumerate(headers, 1):
        cell = _put(sheet, top, column, header, HEAD_FILL)
        cell.font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]
    for index, (values, fill) in enumerate(rows, top + 1):
        for column, value in enumerate(values, 1):
            cell = _put(sheet, index, column, value, fill)
            if isinstance(value, date):
                cell.number_format = "yyyy/mm/dd"
    sheet.freeze_panes = sheet.cell(row=top + 1, column=1)


def render_export_xlsx(data: dict) -> bytes:
    options = data["options"]
    book = Workbook()
    sheet = book.active
    sheet.title = "作業紀錄"
    _put(sheet, 1, 1, "農事作業紀錄表").font = Font(bold=True, size=14)
    info = [("農場名稱", data["farm_name"]), ("經營者", data["operator"]), ("場區", data["field_label"]),
            ("面積", data["area_label"]), ("品項／品種", data["crop_label"]), ("紀錄期間", data["period_label"]),
            ("製表日期", data["generated_label"]), ("欄位規則", data["rule_version"])]
    for index, (label, value) in enumerate(info, 2):
        _put(sheet, index, 1, label).font = Font(bold=True)
        _put(sheet, index, 2, value)
    headers = ["日期", "時間", "作業項目", "場區", "品項", "面積", "資材品名", "供應者", "批號", "稀釋倍數", "用量", "用量單位",
               "產量", "產量單位", "去向", "病蟲害", "人力／工時", "作業說明", "備註", "安全採收期（天）", "安全採收期依據",
               "依據診斷書", "記錄者", "照片張數", "缺漏欄位"]
    rows = []
    for row in data["rows"]:
        log = row["log"]
        spray = log["kind"] == "spray"
        rows.append(([
            log["done_on"], log["done_time"] or "", row["kind_label"], log["field_text"], log["crop_text"], log["area_text"],
            log["material_name"], log["material_supplier"], log["material_lot"], log["dilution"], log["amount_value"],
            log["amount_unit"], log["yield_value"], log["yield_unit"], log["yield_destination"], log["pest_text"],
            log["labor_text"], row["description"], log["note"],
            (log["phi_days"] if log["phi_days"] is not None else "未載明") if spray else "",
            (PHI_SOURCES.get(log["phi_source"]) or "未帶入") if spray else "",
            log.get("case_no") or "", log.get("created_by_name") or "", row["photo_count"], "、".join(row["missing"]),
        ], MISS_FILL if row["miss"] else None))
    _table(sheet, 11, headers, rows, [11, 7, 9, 12, 10, 10, 20, 14, 12, 9, 8, 8, 8, 8, 14, 12, 12, 36, 24, 10, 10, 12, 10, 8, 18])

    harvest = book.create_sheet("採收與產量")
    harvest_rows = [([h["done_on"], h["field_text"], h["item"], h["yield_value"], h["yield_unit"], h["destination"]], None)
                    for h in data["harvests"]]
    harvest_rows += [(["", "", "本期合計", total["total"], total["unit"], ""], None) for total in data["harvest_totals"]]
    if data["unparsed_yields"]:
        harvest_rows.append((["", "", f"另有 {data['unparsed_yields']} 筆產量不是純數字，未列入合計", "", "", ""], None))
    _table(harvest, 1, ["日期", "場區", "品項／批次", "產量", "單位", "去向"], harvest_rows, [11, 14, 24, 10, 8, 24])

    if options.phi:
        check = book.create_sheet("安全採收期核對")
        check_rows = [([entry["harvest_on"], entry["field_text"], entry["crop_text"],
                        entry["spray"]["done_on"] if entry["spray"] else "", entry["spray"]["material_name"] if entry["spray"] else "",
                        entry["spray"]["phi_days"] if entry["spray"] and entry["spray"]["phi_days"] is not None else "",
                        entry["days_since_last"] if entry["days_since_last"] is not None else "", entry["result_label"],
                        entry["message"]], CRIT_FILL if entry["status"] == "violation" else None) for entry in data["phi"]]
        _table(check, 1, ["採收日", "場區", "品項", "施藥日", "藥劑", "安全採收期（天）", "距上次施藥（天）", "結果", "說明"],
               check_rows, [11, 14, 12, 11, 22, 10, 10, 14, 48])

    if options.photos:
        photos = book.create_sheet("照片清單")
        headers = ["附錄編號", "作業日期", "作業項目", "場區", "時間來源", "時間"] + (["緯度", "經度"] if options.gps else []) + ["說明"]
        photo_rows = []
        for item in data["appendix"]:
            values = [item["number"], item["log_date"], item["kind_label"], item["field_text"],
                      "相機拍攝時間" if item["taken_at_source"] == "exif" else "上傳時間（無拍攝資訊）",
                      roc_datetime(item["taken_at"])]
            if options.gps:
                values += [item["lat"] if item["lat"] is not None else "", item["lng"] if item["lng"] is not None else ""]
            photo_rows.append((values + [item["caption"]], None))
        _table(photos, 1, headers, photo_rows, [9, 10, 9, 14, 18, 16] + ([11, 11] if options.gps else []) + [24])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
