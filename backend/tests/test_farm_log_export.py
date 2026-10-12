"""農事紀錄表匯出的測試。

純函式（安全採收期核對、組表、Excel、附錄縮圖）直接跑；API 共用農事日誌的離線 fixture；
PDF 和診斷書一樣用真的 Chromium 產出後讀回文字與字型。資料一律用試驗值，不抄設計稿上的藥名與數字。
"""
import base64
import io
import re
import unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pdfplumber
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import update

from app.core.config import settings
from app.models.farm_log import FarmWorkLog
from app.services.consult_report import page_footer
from app.services.farm_log import COMPLIANCE_RULE_VERSION
from app.services.farm_log_export import (ExportOptions, appendix_image, build_export, phi_checks, render_export_html,
                                          render_export_xlsx)
from tests.test_farm_logs import START, URL, spray, system, upload  # noqa: F401  共用同一套離線 API fixture

MEDIA = Path(__file__).parent / "fixtures" / "media"
EXPORT = "/api/v1/planning/farmers/farmer/logs/export"
PERIOD = {"from": "2030-03-09", "to": "2030-03-10"}
TEXT_KEYS = ("field_text", "crop_text", "title", "material_name", "material_supplier", "material_lot", "dilution",
             "amount_value", "amount_unit", "area_text", "yield_value", "yield_unit", "yield_destination", "labor_text",
             "pest_text", "note")


def entry(kind="spray", done_on=date(2030, 3, 9), **changes):
    base = {key: "" for key in TEXT_KEYS}
    return base | dict(id=f"{kind}-{done_on}", kind=kind, done_on=done_on, done_time=None, plan_id=None,
                       created_at=datetime(2030, 3, 9, tzinfo=timezone.utc), field_text="試驗田甲",
                       compliance={"rule_version": COMPLIANCE_RULE_VERSION, "missing": []}, phi_days=None,
                       phi_source=None, case_no="", created_by_name="試驗農民", photos=[]) | changes


def sprayed(day, phi, **changes):
    return dict(id=f"s-{day}-{phi}", done_on=day, field_text="試驗田甲", plan_id=None, material_name="試驗藥劑",
                phi_days=phi, phi_source="registration") | changes


def photo(source, lat=None):
    return dict(id=f"p-{source}-{lat}", path=None, taken_at=datetime(2030, 3, 9, 1, 30, tzinfo=timezone.utc),
                taken_at_source=source, lat=lat, lng=120.5 if lat is not None else None, caption="")


def export(logs, sprays=(), options=ExportOptions(), field=None):
    return build_export(farm_name="試驗農場", operator="試驗農民", start=date(2030, 3, 1), end=date(2030, 3, 31), field=field,
                        logs=list(logs), sprays=list(sprays), plans={}, options=options, generated_on=date(2030, 4, 1))


# --------------------------------------------------------------------------
# 安全採收期核對
# --------------------------------------------------------------------------

def test_phi_uses_the_most_restrictive_spray_not_the_latest():
    harvest = entry("harvest", date(2030, 3, 9))
    result = phi_checks([harvest], [sprayed(date(2030, 3, 1), 10), sprayed(date(2030, 3, 7), 1)])[0]
    assert result["status"] == "violation" and result["spray"]["done_on"] == date(2030, 3, 1)
    assert result["safe_from"] == date(2030, 3, 11) and result["days_since_last"] == 2
    assert "未達安全採收期 10 天" in result["message"] and result["result_label"] == "未達安全採收期"


def test_phi_scope_boundaries_and_unknown_never_reads_as_safe():
    harvest = entry("harvest", date(2030, 3, 20), plan_id="plan")

    def status(*sprays):
        return phi_checks([harvest], list(sprays))[0]["status"]

    assert status(sprayed(date(2030, 3, 10), 10)) == "ok"            # 剛好滿安全採收期天數就可以採
    assert status(sprayed(date(2030, 3, 11), 10)) == "violation"
    assert status(sprayed(date(2030, 3, 20), 0)) == "ok"
    # 有一次施藥沒有依據，就不能說符合；但確定違反時先講違反。
    assert status(sprayed(date(2030, 3, 10), 3), sprayed(date(2030, 3, 15), None)) == "unknown"
    assert status(sprayed(date(2030, 3, 18), 5), sprayed(date(2030, 3, 15), None)) == "violation"
    assert status(sprayed(date(2030, 3, 19), 5, field_text="試驗田乙")) == "no_spray"   # 別的田區
    assert status(sprayed(date(2030, 3, 19), 5, plan_id="other")) == "no_spray"         # 同田區的上一期作物
    assert status(sprayed(date(2030, 3, 21), 5)) == "no_spray"                          # 採收之後才噴
    assert status(sprayed(date(2029, 3, 1), 400)) == "no_spray"                         # 超過回溯期
    unknown = phi_checks([harvest], [sprayed(date(2030, 3, 15), None, material_name="手寫藥名")])[0]
    assert "手寫藥名" in unknown["message"] and "無法判斷" in unknown["message"]


# --------------------------------------------------------------------------
# 組表
# --------------------------------------------------------------------------

def test_rows_recheck_compliance_number_photos_and_mark_upload_time():
    data = export([
        entry("fertilize", date(2030, 3, 10), material_name="試驗肥料", amount_value="25", amount_unit="公斤",
              area_text="0.4 公頃", compliance={"rule_version": "舊版", "missing": []}, photos=[photo("upload")]),
        entry("spray", date(2030, 3, 9), done_time="06:30", title="施藥 · 試驗藥劑 1000 倍", material_name="試驗藥劑",
              material_supplier="試驗行", material_lot="L-1", dilution="1000", amount_value="0.3", amount_unit="公升",
              area_text="0.4 公頃", pest_text="試驗病害", phi_days=4, phi_source="registration",
              photos=[photo("exif", 22.5), photo("exif", 22.6)]),
        entry("harvest", date(2030, 3, 10), crop_text="試驗作物甲", area_text="0.4 公頃", yield_value="1,200",
              yield_unit="公斤", yield_destination="試驗市場", field_text="試驗田乙"),
        entry("harvest", date(2030, 3, 10), id="h2", crop_text="試驗作物甲", area_text="0.4 公頃", yield_value="約兩籃",
              yield_unit="籃", labor_text="2 人"),
    ], options=ExportOptions(gps=False))
    spray_row, fertilize_row, basket_row, harvest_row = data["rows"]
    assert [row["kind_label"] for row in data["rows"]] == ["施藥", "施肥", "採收", "採收"]
    assert spray_row["material"] == "試驗藥劑／試驗行／L-1" and spray_row["photos_label"] == "1–2"
    assert spray_row["description"] == "試驗病害防治；稀釋 1000 倍；安全採收期 4 天（3/13 起可採）"
    # 存檔時的舊規則說齊全，匯出依目前規則重新提醒，但不改寫紀錄本身。
    assert fertilize_row["miss"] and fertilize_row["material_pending"] == "／（供應者、批號待補）"
    assert fertilize_row["photos_label"] == "3" and data["warnings"]["rule_changed"] == 1
    assert basket_row["amount"] == "約兩籃 籃" and basket_row["description"] == "人力 2 人"
    assert harvest_row["amount"] == "1,200 公斤" and "出貨 試驗市場" in harvest_row["description"]
    assert data["appendix"][0]["time_text"] == "拍攝時間 119/03/09 09:30"
    assert data["appendix"][2]["time_text"] == "上傳時間（無拍攝資訊）119/03/09 09:30"
    assert data["show_field_column"] and data["field_label"] == "試驗田甲、試驗田乙"
    assert [t["text"] for t in data["harvest_totals"]] == ["1,200 公斤"] and data["unparsed_yields"] == 1
    assert data["period_label"] == "119 年 3 月 1 日 至 3 月 31 日"
    assert [item["missing"] for item in data["warnings"]["incomplete"]] == [["供應者", "批號"]]


def test_record_sheet_html_escapes_input_and_shows_phi_and_pending_fields():
    data = export([
        entry("spray", date(2030, 3, 8), material_name="<script>試驗</script>", case_no="T-9", phi_days=5,
              phi_source="prescription", amount_value="1", amount_unit="公升", area_text="0.1 公頃"),
        entry("harvest", date(2030, 3, 10), crop_text="試驗作物甲", area_text="0.1 公頃", yield_value="3", yield_unit="公斤"),
    ], sprays=[sprayed(date(2030, 3, 8), 5)], options=ExportOptions(photos=False))
    html = render_export_html(data, "data:font/otf;base64,")
    assert "<script>" not in html and "&lt;script&gt;" in html
    for expected in ("（供應者、批號待補）", "依診斷書 T-9", "未達安全採收期 5 天", "距上次施藥 2 天", "本次匯出未附照片附錄"):
        assert expected in html, expected
    assert "附錄：照片" not in html


def test_xlsx_keeps_formula_like_text_as_text_and_adds_counts():
    data = export([entry("spray", material_name='=HYPERLINK("http://example.invalid")', note="+試驗備註",
                         photos=[photo("exif", 22.5)])])
    book = load_workbook(io.BytesIO(render_export_xlsx(data)))
    assert book.sheetnames == ["作業紀錄", "採收與產量", "安全採收期核對", "照片清單"]
    sheet = book["作業紀錄"]
    headers = [cell.value for cell in sheet[11]]
    row = dict(zip(headers, sheet[12]))
    assert row["資材品名"].value.startswith("=HYPERLINK") and row["資材品名"].data_type == "s"
    assert row["照片張數"].value == 1 and "供應者" in row["缺漏欄位"].value
    assert row["安全採收期（天）"].value == "未載明" and row["安全採收期依據"].value == "未帶入"
    assert [cell.value for cell in book["照片清單"][1]][6:8] == ["緯度", "經度"]


def test_appendix_image_follows_exif_orientation_and_tolerates_missing_files(tmp_path):
    url = appendix_image(MEDIA / "exif_gps_rotated.jpg")
    image = Image.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1])))
    assert image.format == "JPEG" and image.size == (40, 64)  # 64x40、方向 6：轉正後直立
    large = Image.open(io.BytesIO(base64.b64decode(appendix_image(MEDIA / "large.jpg").split(",", 1)[1])))
    assert max(large.size) <= 1000
    assert appendix_image(tmp_path / "missing.jpg") is None and appendix_image(None) is None


# --------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------

def seed_period(clients, clock):
    farmer = clients["farmer"]
    created = {}

    def post(name, payload):
        clock["now"] += timedelta(minutes=1)  # 同日紀錄依建立時間排序，測試要能預期順序
        response = farmer.post(URL, json=payload)
        assert response.status_code == 201, response.text
        created[name] = response.json()

    # 期間前一天的施藥：核對要往回找，只看匯出期間會漏掉最要緊的那一筆。
    post("before", spray(done_on="2030-03-08", registration_id=1))
    post("rx", spray(prescription_item_id="rx", material_name="", note='=HYPERLINK("http://example.invalid")'))
    post("fertilize", dict(kind="fertilize", done_on="2030-03-09", field_text="試驗田甲", material_name="試驗肥料",
                           amount_value="2", amount_unit="公斤", area_text="0.4 公頃"))
    post("harvest", dict(kind="harvest", done_on="2030-03-10", field_text="試驗田甲", crop_text="試驗作物甲",
                         area_text="0.4 公頃", yield_value="12", yield_unit="公斤", yield_destination="試驗市場"))
    for name in ("exif_gps_rotated.jpg", "no_exif.png"):
        assert upload(farmer, created["harvest"]["id"], name).status_code == 201
    return created


def test_export_requires_farm_access_and_a_valid_period(system):
    app, _, clients, _ = system
    params = PERIOD | {"format": "xlsx"}
    assert TestClient(app).get(EXPORT, params=params).status_code == 401
    for outsider in ("stranger", "unlinked"):
        assert clients[outsider].get(EXPORT, params=params).status_code == 404
        assert clients[outsider].get(EXPORT + "/preview", params=PERIOD).status_code == 404
    assert clients["doc"].get(EXPORT, params=params).status_code == 200  # 已連結的植物醫生可以替農場匯出
    farmer = clients["farmer"]
    assert farmer.get(EXPORT, params={"from": "2030-03-10", "to": "2030-03-01"}).status_code == 422
    too_long = farmer.get(EXPORT, params={"from": "2028-01-01", "to": "2030-03-01", "format": "xlsx"})
    assert too_long.status_code == 422 and "分段" in too_long.json()["detail"]
    assert farmer.get(EXPORT, params=PERIOD | {"format": "docx"}).status_code == 422
    assert farmer.get(EXPORT, params={"to": "2030-03-01"}).status_code == 422


def test_preview_rechecks_compliance_and_flags_phi_from_before_the_period(system):
    _, factory, clients, clock = system
    created = seed_period(clients, clock)
    with factory() as db:
        # 模擬舊版規則存下的判定：當時說齊全，匯出時要依目前規則重新提醒。
        db.execute(update(FarmWorkLog).where(FarmWorkLog.id == created["fertilize"]["id"])
                   .values(compliance={"rule_version": "舊版", "required": [], "missing": []}))
        db.commit()
    body = clients["farmer"].get(EXPORT + "/preview", params=PERIOD).json()
    assert (body["total"], body["photo_count"], body["too_many_photos"]) == (3, 2, False)
    assert body["rule_changed"] == 1 and body["rule_version"] == COMPLIANCE_RULE_VERSION
    fertilize = next(item for item in body["incomplete"] if item["id"] == created["fertilize"]["id"])
    assert fertilize["missing"] == ["供應者", "批號"]
    assert body["phi"]["checked"] == 1 and body["phi"]["unknown"] == []
    violation = body["phi"]["violations"][0]
    assert violation["harvest_on"] == "2030-03-10" and "未達安全採收期 4 天" in violation["message"]
    field_only = clients["farmer"].get(EXPORT + "/preview", params=PERIOD | {"field": "試驗田乙"}).json()
    assert field_only["total"] == 0 and field_only["phi"]["checked"] == 0


def test_xlsx_download_matches_log_order_and_sources(system):
    _, _, clients, clock = system
    seed_period(clients, clock)
    response = clients["farmer"].get(EXPORT, params=PERIOD | {"format": "xlsx", "gps": "0"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["cache-control"] == "private, no-store"
    disposition = response.headers["content-disposition"]
    assert 'filename="farm-log-20300309-20300310.xlsx"' in disposition and "filename*=UTF-8''" in disposition
    book = load_workbook(io.BytesIO(response.content))
    sheet = book["作業紀錄"]
    headers = [cell.value for cell in sheet[11]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=12, values_only=True)]
    assert [row["作業項目"] for row in rows] == ["施藥", "施肥", "採收"]
    assert (rows[0]["依據診斷書"], rows[0]["安全採收期依據"], rows[0]["安全採收期（天）"]) == ("T-1", "處方", "未載明")
    assert rows[0]["備註"].startswith("=HYPERLINK") and sheet.cell(row=12, column=headers.index("備註") + 1).data_type == "s"
    assert rows[2]["照片張數"] == 2 and rows[1]["缺漏欄位"] == "供應者、批號"
    assert "緯度" not in [cell.value for cell in book["照片清單"][1]]
    assert list(book["安全採收期核對"].iter_rows(min_row=2, values_only=True))[0][7] == "未達安全採收期"


FOOTER_PT = 15 / 25.4 * 72  # 版面下邊界 15mm；頁碼印在這裡


def pdf_text(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    with pdfplumber.open(io.BytesIO(response.content)) as pdf:
        body = [char for page in pdf.pages for char in page.chars if char["top"] < page.height - FOOTER_PT]
        # 內文字型一律是隨程式附帶的思源宋體；主機沒有中文字型時也不會變成方塊。頁尾由 Chromium 用系統字型畫。
        assert body and all("SourceHanSerif" in char["fontname"] for char in body)
        footers = ["".join(char["text"] for char in page.chars if char["top"] >= page.height - FOOTER_PT).replace(" ", "")
                   for page in pdf.pages]
        pages = len(pdf.pages)
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    return pages, re.sub(r"\s+", "", unicodedata.normalize("NFKC", text)), footers


def test_page_footer_never_omits_page_numbers():
    chinese = page_footer(True)
    assert "第 " in chinese and "AgriPageNumber" in chinese
    fallback = page_footer(False)
    assert "第" not in fallback and "AgriPageNumber" not in fallback
    for footer in (chinese, fallback):
        assert 'class="pageNumber"' in footer and 'class="totalPages"' in footer


def test_pdf_record_sheet_with_photo_appendix_and_gps_toggle(system, monkeypatch, tmp_path):
    _, _, clients, clock = system
    seed_period(clients, clock)
    # 本機沒有安裝頁碼字型：退回只印數字，但每一頁都要有「第幾頁／共幾頁」。
    monkeypatch.setattr(settings, "pdf_page_font_file", tmp_path / "missing.otf")
    pages, text, footers = pdf_text(clients["farmer"].get(EXPORT, params=PERIOD))
    assert pages >= 2 and footers == [f"{number}/{pages}" for number in range(1, pages + 1)]
    for expected in ("農事作業紀錄表", "附錄1", "附錄2", "拍攝時間119/03/09", "上傳時間(無拍攝資訊)", "GPS22.75847"):
        assert expected in text, expected
    _, without_gps, _ = pdf_text(clients["farmer"].get(EXPORT, params=PERIOD | {"gps": "0"}))
    assert "附錄1" in without_gps and "GPS22.75847" not in without_gps


def test_photo_limit_only_blocks_the_pdf_appendix(system, monkeypatch):
    _, _, clients, clock = system
    seed_period(clients, clock)
    monkeypatch.setattr(settings, "media_export_max_photos", 1)
    blocked = clients["farmer"].get(EXPORT, params=PERIOD)
    assert blocked.status_code == 422 and "上限 1 張" in blocked.json()["detail"]
    assert clients["farmer"].get(EXPORT + "/preview", params=PERIOD).json()["too_many_photos"] is True
    assert clients["farmer"].get(EXPORT, params=PERIOD | {"format": "xlsx"}).status_code == 200
